"""
知识库服务：文档解析 + 上传/列表/删除
从 upload.py 中分离出来，让路由层只处理 HTTP
"""
import io
from fastapi import HTTPException, status
from sqlalchemy.orm import Session
from backend.models import Company, Document
from backend.schemas import DocumentInfo

# 允许的文件类型
ALLOWED_TYPES = {
    "text/plain": "txt",
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
}
MAX_FILE_SIZE = 10 * 1024 * 1024  # 10MB


def sanitize_text(text: str) -> str:
    """清除非法 UTF-8 代理字符（PDF 提取时容易出现）"""
    return text.encode("utf-8", errors="replace").decode("utf-8")


def parse_txt(content: bytes) -> str:
    """解析 TXT 文件"""
    return sanitize_text(content.decode("utf-8", errors="replace"))


def parse_pdf(content: bytes) -> str:
    """解析 PDF 文件"""
    from PyPDF2 import PdfReader
    reader = PdfReader(io.BytesIO(content))
    text_parts = []
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text_parts.append(page_text)
    return sanitize_text("\n".join(text_parts))


def parse_docx(content: bytes) -> str:
    """解析 Word 文档"""
    from docx import Document as DocxDocument
    doc = DocxDocument(io.BytesIO(content))
    text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    return sanitize_text(text)


def parse_file(filename: str, content: bytes) -> tuple[str, str]:
    """根据文件类型解析，返回 (纯文本, 文件类型)"""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if ext == "txt":
        return parse_txt(content), "txt"
    elif ext == "pdf":
        return parse_pdf(content), "pdf"
    elif ext in ("docx", "doc"):
        return parse_docx(content), "docx"
    else:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"不支持的文件类型：.{ext}。支持：txt, pdf, docx",
        )


def upload_document(
    company: Company,
    filename: str,
    file_content: bytes,
    db: Session,
) -> dict:
    """上传知识库文档：校验 → 解析 → 入库 → 触发向量化"""
    if len(file_content) > MAX_FILE_SIZE:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"文件过大（最大 {MAX_FILE_SIZE // 1024 // 1024}MB）",
        )

    text, file_type = parse_file(filename, file_content)

    if not text.strip():
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="文件内容为空，无法解析",
        )

    doc = Document(
        company_id=company.id,
        filename=filename,
        file_type=file_type,
        content=text,
    )
    db.add(doc)
    db.commit()
    db.refresh(doc)

    # 自动分块 + 向量化
    from backend.embedding import process_document
    try:
        chunk_count = process_document(doc.id, company.id, text)
        doc.chunk_count = chunk_count
        db.commit()
    except Exception as e:
        print(f"[WARN] 向量化失败：{e}")

    return {
        "message": "上传成功",
        "document": DocumentInfo(
            id=doc.id,
            filename=doc.filename,
            file_type=doc.file_type,
            chunk_count=doc.chunk_count,
            uploaded_at=doc.uploaded_at.isoformat(),
        ),
    }


def list_documents(company: Company, db: Session) -> list[DocumentInfo]:
    """获取知识库文档列表"""
    docs = (
        db.query(Document)
        .filter(Document.company_id == company.id)
        .order_by(Document.uploaded_at.desc())
        .all()
    )
    return [
        DocumentInfo(
            id=d.id,
            filename=d.filename,
            file_type=d.file_type,
            chunk_count=d.chunk_count,
            uploaded_at=d.uploaded_at.isoformat(),
        )
        for d in docs
    ]


def delete_document(document_id: int, company: Company, db: Session) -> dict:
    """删除知识库文档及关联向量"""
    doc = (
        db.query(Document)
        .filter(Document.id == document_id, Document.company_id == company.id)
        .first()
    )
    if not doc:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="文档不存在")

    from backend.vector_store import vector_store
    vector_store.delete_by_document(doc.id, db=db)

    db.delete(doc)
    db.commit()
    return {"message": "删除成功"}
